"""Convert eval JSONL datasets ↔ Excel (.xlsx) for human editing.

Lists are stored as `` | ``-separated cells. Nested dicts use dotted keys
(e.g. ``input.source_text``, ``input.context.channel``, ``expect.must_contain_any``).
Empty cells become JSON ``null`` on import for known nullable string fields.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

LIST_SEPARATOR = " | "

# Fields that must round-trip as JSON arrays (not plain strings).
_LIST_SUFFIXES = (
    "must_contain_any",
    "required_fields",
    "required_item_types_any",
    "surface_must_contain_any",
    "candidate_types",
)

# input.* paths that are strings and should be null when blank (not "").
_NULLABLE_STRING_PATHS = frozenset(
    {
        "input.incoming_text",
        "input.target_language",
        "input.locale_variant",
        "input.channel",
        "input.source_language",
        "input.spanish_text",
        "input.source_type",
    }
)


def _is_list_key(dotted: str) -> bool:
    leaf = dotted.rsplit(".", 1)[-1]
    return leaf in _LIST_SUFFIXES or dotted.endswith(_LIST_SUFFIXES)


def flatten_case(case: dict[str, Any]) -> dict[str, Any]:
    """Flatten one JSONL case into a single-level dict for a spreadsheet row."""
    row: dict[str, Any] = {}

    def walk(prefix: str, value: Any) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                if key.startswith("_"):
                    continue
                walk(f"{prefix}.{key}" if prefix else key, child)
            return
        if isinstance(value, list):
            row[prefix] = LIST_SEPARATOR.join(str(item) for item in value)
            return
        if value is None:
            row[prefix] = ""
            return
        if isinstance(value, bool):
            row[prefix] = "true" if value else "false"
            return
        row[prefix] = value

    for key in ("id", "suite"):
        if key in case:
            row[key] = case[key]
    if "input" in case:
        walk("input", case["input"])
    if "expect" in case:
        walk("expect", case["expect"])
    # Preserve unknown top-level keys (except loader metadata).
    for key, value in case.items():
        if key in {"id", "suite", "input", "expect"} or key.startswith("_"):
            continue
        walk(key, value)
    return row


def _parse_cell(dotted: str, raw: Any) -> Any:
    leaf = dotted.rsplit(".", 1)[-1]

    if isinstance(raw, bool):
        return raw
    if isinstance(raw, int | float) and not isinstance(raw, bool):
        if leaf.startswith("min_") or leaf.endswith(("_chars", "_candidates")):
            return int(raw) if float(raw).is_integer() else raw
        if isinstance(raw, float) and raw.is_integer():
            return int(raw)
        return raw

    text = "" if raw is None else str(raw).strip()

    if text == "":
        if dotted in _NULLABLE_STRING_PATHS:
            return None
        if _is_list_key(dotted):
            return []
        if leaf.startswith("forbid_"):
            return False
        if leaf.startswith("min_"):
            return 0
        return None if dotted.startswith("input.") else ""

    lower = text.lower()
    if leaf.startswith("forbid_"):
        if lower in {"true", "1", "yes", "y"}:
            return True
        if lower in {"false", "0", "no", "n"}:
            return False

    if _is_list_key(dotted):
        return [p.strip() for p in text.split("|") if p.strip()]

    if leaf.startswith("min_"):
        try:
            return int(float(text))
        except ValueError:
            return text

    if lower in {"true", "false"} and dotted.startswith("expect."):
        return lower == "true"

    return text


def unflatten_row(row: dict[str, Any]) -> dict[str, Any]:
    """Rebuild a JSONL case dict from a flat spreadsheet row."""
    tree: dict[str, Any] = {}

    def assign(path: str, value: Any) -> None:
        parts = path.split(".")
        cursor = tree
        for part in parts[:-1]:
            nxt = cursor.get(part)
            if not isinstance(nxt, dict):
                nxt = {}
                cursor[part] = nxt
            cursor = nxt
        cursor[parts[-1]] = value

    for key, raw in row.items():
        if key is None:
            continue
        dotted = str(key).strip()
        if not dotted or dotted.startswith("_"):
            continue
        assign(dotted, _parse_cell(dotted, raw))

    case: dict[str, Any] = {}
    if "id" in tree:
        case["id"] = tree.pop("id")
    if "suite" in tree:
        case["suite"] = tree.pop("suite")
    if "input" in tree:
        case["input"] = tree.pop("input")
        ctx = case["input"].get("context")
        if isinstance(ctx, dict):
            case["input"]["context"] = {k: v for k, v in ctx.items() if v is not None}
            if not case["input"]["context"]:
                case["input"].pop("context", None)
    if "expect" in tree:
        case["expect"] = tree.pop("expect")
    case.update(tree)
    return case


def cases_from_jsonl(path: Path) -> list[dict[str, Any]]:
    cases: list[dict[str, Any]] = []
    for line_no, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        text = line.strip()
        if not text or text.startswith("#"):
            continue
        try:
            cases.append(json.loads(text))
        except json.JSONDecodeError as exc:
            raise ValueError(f"Invalid JSONL {path.name}:{line_no}") from exc
    return cases


def write_jsonl(cases: list[dict[str, Any]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [json.dumps(case, ensure_ascii=False) for case in cases]
    path.write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8")


def export_jsonl_to_xlsx(jsonl_path: Path, xlsx_path: Path) -> int:
    from openpyxl import Workbook

    cases = cases_from_jsonl(jsonl_path)
    flat_rows = [flatten_case(case) for case in cases]
    headers: list[str] = []
    seen: set[str] = set()
    preferred = [
        "id",
        "suite",
        "input.source_language",
        "input.source_text",
        "input.incoming_text",
        "input.target_language",
        "input.locale_variant",
        "input.channel",
        "input.context.channel",
        "input.context.user_level",
        "input.context.target_locale",
        "input.spanish_text",
        "input.source_type",
        "input.candidate_types",
        "expect.must_contain_any",
        "expect.forbid_cyrillic_in_reply",
        "expect.forbid_latin_letters",
        "expect.forbid_cyrillic",
        "expect.min_reply_chars",
        "expect.min_chars",
        "expect.min_candidates",
        "expect.translation_field",
        "expect.required_item_types_any",
        "expect.surface_must_contain_any",
    ]
    for key in preferred:
        if any(key in row for row in flat_rows) and key not in seen:
            headers.append(key)
            seen.add(key)
    extra = sorted({key for row in flat_rows for key in row if key not in seen})
    headers.extend(extra)

    wb = Workbook()
    ws = wb.active
    ws.title = "cases"
    ws.append(headers)
    for row in flat_rows:
        ws.append([row.get(h, "") for h in headers])
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    return len(cases)


def import_xlsx_to_jsonl(xlsx_path: Path, jsonl_path: Path) -> int:
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise ValueError(f"Empty workbook: {xlsx_path}") from exc
    headers = [str(h).strip() if h is not None else "" for h in header_row]
    if not headers or "id" not in headers:
        raise ValueError("First row must include an 'id' column")

    cases: list[dict[str, Any]] = []
    for excel_row in rows_iter:
        mapping = {headers[i]: excel_row[i] if i < len(excel_row) else None for i in range(len(headers))}
        if mapping.get("id") in (None, ""):
            continue
        case = unflatten_row(mapping)
        if "suite" not in case or not case["suite"]:
            raise ValueError(f"Row id={case.get('id')!r} missing suite")
        if "input" not in case:
            raise ValueError(f"Row id={case.get('id')!r} missing input.* columns")
        cases.append(case)

    write_jsonl(cases, jsonl_path)
    return len(cases)
